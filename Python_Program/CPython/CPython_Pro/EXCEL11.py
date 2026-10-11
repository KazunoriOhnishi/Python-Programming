import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# 1. 新しいワークブック（Excelファイル）を作成
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "売上表"

# 2. 表のデータを定義
data = [
    ["商品名", "単価", "数量", "売上金額"],
    ["ノートPC", 120000, 5, "=B2*C2"],
    ["液晶モニター", 25000, 10, "=B3*C3"],
    ["ワイヤレスマウス", 3500, 30, "=B4*C4"],
    ["合計", "", "", "=SUM(D2:D4)"]
]

# 3. データをセルに書き込み
for row in data:
    ws.append(row)

# 4. 表のデザイン（装飾）を設定
# スタイルの準備
header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid") # 青色
header_font = Font(name="Meiryo", size=11, bold=True, color="FFFFFF") # 白文字・太字
data_font = Font(name="Meiryo", size=11)
thin_side = Side(border_style="thin", color="000000")
thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

# 見出し行（1行目）への装飾適用
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")
    cell.border = thin_border

# データ行への装飾適用
for row in ws.iter_rows(min_row=2, max_row=5, min_col=1, max_col=4):
    for cell in row:
        cell.font = data_font
        cell.border = thin_border
        # 数値の列は右寄せ、カンマ区切りフォーマットにする
        if cell.column in [2, 3, 4] and cell.value != "":
            cell.alignment = Alignment(horizontal="right")
            if not str(cell.value).startswith("="): # 数式以外の数値
                cell.number_format = '#,##0'

# 5. ファイルを保存
wb.save("売上データ.xlsx")
print("Excelファイル『売上データ.xlsx』を作成しました。")
